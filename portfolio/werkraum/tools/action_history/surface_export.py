"""Export a captured flextrawurst Surface as a navigable document.

The Surface is a tabbed application, not a chat transcript.  This module keeps
the tab inventory and turns every captured tab into a stable document section.
The HTML export uses normal fragment links; PDF and DOCX receive internal
hyperlinks so the tab index remains navigable after JavaScript is gone.
"""

from __future__ import annotations

import base64
import hashlib
import html
import io
import json
import re
from pathlib import Path
from typing import Any

try:
    from .office_export import FORMATS, OfficeArtifact, OfficeExportError, _pdf_fonts
except ImportError:  # direct CLI execution from tools/action_history
    from office_export import FORMATS, OfficeArtifact, OfficeExportError, _pdf_fonts


def _slug(value: Any, fallback: str = "tab") -> str:
    clean = re.sub(r"[^a-zA-Z0-9_-]+", "-", str(value or "").strip().lower()).strip("-")
    return (clean[:80] or fallback)


def _tabs(payload: dict[str, Any]) -> list[dict[str, Any]]:
    tabs = payload.get("tabs") if isinstance(payload.get("tabs"), list) else []
    result: list[dict[str, Any]] = []
    seen: set[str] = set()
    for index, raw in enumerate(tabs, start=1):
        if not isinstance(raw, dict):
            continue
        tab_id = str(raw.get("id") or raw.get("view_id") or f"tab-{index}").strip()
        if not tab_id or tab_id in seen:
            continue
        seen.add(tab_id)
        label = str(raw.get("label") or tab_id).strip()[:240]
        anchor = f"surface-tab-{_slug(tab_id, f'tab-{index}')}-{index}"
        active_display = str(raw.get("active_display") or "")
        scroll_stalled = bool(raw.get("scroll_stalled"))
        capture_status = "sichtbar-vollständig" if active_display not in {"", "none"} and not scroll_stalled else "nicht-aktiv-bestätigt"
        result.append({
            **raw,
            "id": tab_id,
            "label": label or tab_id,
            "capture_status": capture_status,
            "anchor": anchor,
            "content": str(raw.get("content") or "").strip(),
            "links": [item for item in raw.get("links", []) if isinstance(item, dict)] if isinstance(raw.get("links"), list) else [],
        })
    return result


def render_surface_markdown(payload: dict[str, Any]) -> str:
    title = str(payload.get("title") or "flextrawurst Surface").strip()
    source_url = str(payload.get("source_url") or "").strip()
    captured = str(payload.get("captured_utc") or "").strip()
    tabs = _tabs(payload)
    lines = [f"# {title}", "", f"Quelle: {source_url}" if source_url else "Quelle: unbekannt", f"Aufgenommen: {captured}" if captured else "", "", "## Surface-Tabs", ""]
    for tab in tabs:
        lines.append(f"- [{tab['label']} — {tab['capture_status']}](#{tab['anchor']})")
    lines.extend(["", "---", ""])
    for index, tab in enumerate(tabs, start=1):
        lines.extend([f"<a id=\"{tab['anchor']}\"></a>", f"## {index}. {tab['label']} · {tab['capture_status']}", ""])
        if tab["content"]:
            lines.extend([tab["content"], ""])
        else:
            lines.extend(["*(Dieser Tab lieferte im Lauf keinen sichtbaren Text.)*", ""])
        if tab["links"]:
            lines.extend(["Links in diesem Tab:", ""])
            for link in tab["links"]:
                href = str(link.get("href") or "").strip()
                label = str(link.get("text") or href).strip()
                if href:
                    lines.append(f"- [{label}]({href})")
            lines.append("")
        lines.append("[Zur Tab-Übersicht](#surface-tab-index)")
        lines.append("")
    return "\n".join(lines).rstrip() + "\n"


def render_surface_html(payload: dict[str, Any]) -> str:
    title = str(payload.get("title") or "flextrawurst Surface").strip()
    source_url = str(payload.get("source_url") or "").strip()
    captured = str(payload.get("captured_utc") or "").strip()
    tabs = _tabs(payload)
    nav = "".join(f'<li><a href="#{html.escape(tab["anchor"], quote=True)}">{html.escape(tab["label"])}</a> · {html.escape(tab["capture_status"])}</li>' for tab in tabs)
    sections: list[str] = []
    for index, tab in enumerate(tabs, start=1):
        links = "".join(
            f'<li><a href="{html.escape(str(link.get("href") or ""), quote=True)}">{html.escape(str(link.get("text") or link.get("href") or "Link"))}</a></li>'
            for link in tab["links"] if str(link.get("href") or "").strip()
        )
        sections.append(
            f'<section id="{html.escape(tab["anchor"], quote=True)}" data-surface-tab="{html.escape(tab["id"], quote=True)}">'
            f'<h2>{index}. {html.escape(tab["label"])} <small>· {html.escape(tab["capture_status"])}</small></h2>'
            f'<div class="surface-content">{html.escape(tab["content"]).replace(chr(10), "<br>") if tab["content"] else "<em>Kein sichtbarer Text in diesem Tab.</em>"}</div>'
            f'{f"<h3>Links</h3><ul>{links}</ul>" if links else ""}'
            '<p><a href="#surface-tab-index">↑ Zur Tab-Übersicht</a></p></section>'
        )
    return """<!doctype html>
<html lang="de"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>""" + html.escape(title) + """ — Surface-Export</title>
<style>body{font:16px/1.55 system-ui,sans-serif;max-width:1000px;margin:2rem auto;padding:0 1.2rem;color:#202124}header{border-bottom:2px solid #10d8f0;padding-bottom:1rem}nav{background:#f5f8fa;padding:1rem;margin:1.5rem 0}section{border-top:1px solid #d8e0e5;padding:1.5rem 0;scroll-margin-top:1rem}.surface-content{white-space:pre-wrap}a{color:#075ca8}</style></head><body>
<header><h1>""" + html.escape(title) + """ — Surface-Export</h1>""" + (f'<p>Quelle: <a href="{html.escape(source_url, quote=True)}">{html.escape(source_url)}</a></p>' if source_url else "") + (f'<p>Aufgenommen: {html.escape(captured)}</p>' if captured else "") + """</header>
<nav id="surface-tab-index"><h2>Surface-Tabs</h2><ol>""" + nav + """</ol></nav>""" + "".join(sections) + "</body></html>\n"


def _docx_bookmark(paragraph: Any, text: str, anchor: str, bookmark_id: int) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), anchor)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    run = paragraph.add_run(text)
    run.bold = True
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def _docx_internal_link(paragraph: Any, text: str, anchor: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    link = OxmlElement("w:hyperlink")
    link.set(qn("w:anchor"), anchor)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "075CA8")
    props.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(underline)
    run.append(props)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    link.append(run)
    paragraph._p.append(link)


def _docx_external_link(paragraph: Any, text: str, href: str) -> None:
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    rel_id = paragraph.part.relate_to(href, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "075CA8")
    props.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(underline)
    run.append(props)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    link.append(run)
    paragraph._p.append(link)


def _render_surface_docx(payload: dict[str, Any]) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt

    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.65)
    section.left_margin = section.right_margin = Inches(0.72)
    document.styles["Normal"].font.size = Pt(10.5)
    title = document.add_heading(str(payload.get("title") or "flextrawurst Surface"), 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.add_paragraph("Surface-Export · jeder Tab als eigener Abschnitt")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if payload.get("source_url"):
        p = document.add_paragraph("Quelle: ")
        _docx_external_link(p, str(payload["source_url"]), str(payload["source_url"]))
    if payload.get("captured_utc"):
        document.add_paragraph(f"Aufgenommen: {payload['captured_utc']}")
    tabs = _tabs(payload)
    document.add_heading("Surface-Tabs", level=1)
    for tab in tabs:
        p = document.add_paragraph(style="List Number")
        _docx_internal_link(p, f"{tab['label']} — {tab['capture_status']}", tab["anchor"])
    document.add_page_break()
    for index, tab in enumerate(tabs, start=1):
        heading = document.add_paragraph()
        _docx_bookmark(heading, f"{index}. {tab['label']} · {tab['capture_status']}", tab["anchor"], index)
        heading.style = document.styles["Heading 1"]
        if tab["content"]:
            for chunk in tab["content"].split("\n\n"):
                document.add_paragraph(chunk.strip())
        else:
            document.add_paragraph("Kein sichtbarer Text in diesem Tab.").italic = True
        for link in tab["links"]:
            href = str(link.get("href") or "").strip()
            if href:
                p = document.add_paragraph(style="List Bullet")
                _docx_external_link(p, str(link.get("text") or href), href)
        p = document.add_paragraph()
        _docx_internal_link(p, "↑ Zur Tab-Übersicht", "surface-tab-index")
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def _render_surface_pdf(payload: dict[str, Any]) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    regular, bold, _mono = _pdf_fonts()
    out = io.BytesIO()
    document = SimpleDocTemplate(out, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm, bottomMargin=18 * mm)
    sheet = getSampleStyleSheet()
    title_style = ParagraphStyle("SurfaceTitle", parent=sheet["Title"], fontName=bold, fontSize=19, leading=23, alignment=TA_CENTER, spaceAfter=5)
    sub_style = ParagraphStyle("SurfaceSub", parent=sheet["Normal"], fontName=regular, fontSize=9, textColor=colors.HexColor("#555555"), alignment=TA_CENTER, spaceAfter=10)
    h1 = ParagraphStyle("SurfaceH1", parent=sheet["Heading1"], fontName=bold, fontSize=14, leading=17, spaceBefore=10, spaceAfter=6)
    body = ParagraphStyle("SurfaceBody", parent=sheet["BodyText"], fontName=regular, fontSize=9.5, leading=13.5, spaceAfter=7)
    link_style = ParagraphStyle("SurfaceLink", parent=body, textColor=colors.HexColor("#075CA8"))
    story: list[Any] = [Paragraph(html.escape(str(payload.get("title") or "flextrawurst Surface")), title_style), Paragraph("Surface-Export · jeder Tab als eigener Abschnitt", sub_style)]
    if payload.get("source_url"):
        source = html.escape(str(payload["source_url"]), quote=True)
        story.append(Paragraph(f'Quelle: <a href="{source}" color="#075CA8"><u>{source}</u></a>', body))
    if payload.get("captured_utc"):
        story.append(Paragraph(f"Aufgenommen: {html.escape(str(payload['captured_utc']))}", body))
    tabs = _tabs(payload)
    story.append(Paragraph('<a name="surface-tab-index"/>Surface-Tabs', h1))
    for index, tab in enumerate(tabs, start=1):
        story.append(Paragraph(f'{index}. <a href="#{html.escape(tab["anchor"], quote=True)}" color="#075CA8"><u>{html.escape(tab["label"])}</u></a> · {html.escape(tab["capture_status"])}', link_style))
    for index, tab in enumerate(tabs, start=1):
        story.append(Spacer(1, 4))
        story.append(Paragraph(f'<a name="{html.escape(tab["anchor"], quote=True)}"/>{index}. {html.escape(tab["label"])} · {html.escape(tab["capture_status"])}', h1))
        content = html.escape(tab["content"]).replace("\n", "<br/>") if tab["content"] else "<i>Kein sichtbarer Text in diesem Tab.</i>"
        story.append(Paragraph(content, body))
        for link in tab["links"]:
            href = html.escape(str(link.get("href") or ""), quote=True)
            label = html.escape(str(link.get("text") or href))
            if href:
                story.append(Paragraph(f'· <a href="{href}" color="#075CA8"><u>{label}</u></a>', body))
        story.append(Paragraph(f'<a href="#surface-tab-index" color="#075CA8"><u>↑ Zur Tab-Übersicht</u></a>', link_style))

    def footer(canvas: Any, _doc: Any) -> None:
        canvas.saveState()
        canvas.setFont(regular, 7.5)
        canvas.setFillColor(colors.HexColor("#666666"))
        canvas.drawString(18 * mm, 10 * mm, "flextrawurst Surface-Export")
        canvas.drawRightString(A4[0] - 18 * mm, 10 * mm, f"Seite {canvas.getPageNumber()}")
        canvas.restoreState()

    document.build(story, onFirstPage=footer, onLaterPages=footer)
    return out.getvalue()


def build_surface_office_artifact(payload: dict[str, Any], *, format_name: str) -> OfficeArtifact:
    if format_name not in {"pdf", "docx"}:
        raise OfficeExportError("Surface-Export unterstützt pdf oder docx")
    if not _tabs(payload):
        raise OfficeExportError("Surface-Aufnahme enthält keine Tabs")
    content = _render_surface_pdf(payload) if format_name == "pdf" else _render_surface_docx(payload)
    title = _slug(payload.get("title"), "flextrawurst-surface")
    return OfficeArtifact(content=content, filename=f"{title}-surface-tabs.{format_name}", mime_type=FORMATS[format_name], format=format_name, variant="surface_tabs")


def write_surface_artifacts(payload: dict[str, Any], output_dir: str | Path) -> dict[str, Any]:
    folder = Path(output_dir)
    folder.mkdir(parents=True, exist_ok=True)
    markdown = render_surface_markdown(payload).encode("utf-8")
    html_doc = render_surface_html(payload).encode("utf-8")
    artifacts: dict[str, Any] = {}
    for name, content in (("SURFACE.md", markdown), ("SURFACE.html", html_doc)):
        path = folder / name
        path.write_bytes(content)
        artifacts[name] = {"path": str(path), "bytes": len(content), "sha256": hashlib.sha256(content).hexdigest()}
    for format_name in ("pdf", "docx"):
        artifact = build_surface_office_artifact(payload, format_name=format_name)
        path = folder / artifact.filename
        path.write_bytes(artifact.content)
        artifacts[path.name] = {"path": str(path), "bytes": len(artifact.content), "sha256": hashlib.sha256(artifact.content).hexdigest(), "mime_type": artifact.mime_type}
    manifest = {"schema_version": 1, "artifact_kind": "flextrawurst_surface_export", "source_url": payload.get("source_url"), "title": payload.get("title"), "captured_utc": payload.get("captured_utc"), "tab_count": len(_tabs(payload)), "internal_tab_link_count": len(_tabs(payload)) * 2, "artifacts": artifacts}
    (folder / "SURFACE-MANIFEST.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return manifest
